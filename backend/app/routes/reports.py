import io
from datetime import date

from flask import request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..constants import BRANCHES, DENOMINATIONS
from ..models import (
    CellReport,
    Department,
    DepartmentMember,
    DepartmentReport,
    FundContribution,
    HomeChurch,
    InventoryItem,
    Leader,
    Member,
    Offering,
    Pledge,
    Tithe,
)
from ..extensions import db
from ..security import ApiError, apply_scope, audit, login_required, scope_branch
from . import api
from .people import OFFICE

BRANCH_NAME = {i: n for i, _, n in BRANCHES}
GROUP_LABEL = {"sunday_school": "Sunday School", "teens": "Teenagers", "junior_youth": "Junior Youth", "senior_youth": "Senior Youth", "fathers": "Fathers", "mothers": "Mothers"}
METHOD = {"mpesa": "M-Pesa", "bank": "Bank", "cash": "Cash"}
HEADER_FILL = PatternFill("solid", fgColor="2D3191")
TOTAL_FILL = PatternFill("solid", fgColor="F3F0E8")


def _dated(q, model, column="date"):
    col = getattr(model, column)
    a = request.args
    if a.get("from"):
        q = q.filter(col >= a["from"])
    if a.get("to"):
        q = q.filter(col <= a["to"])
    return q


def _period():
    a = request.args
    if not a.get("from") and not a.get("to"):
        return "All dates"
    return f"{a.get('from') or 'start'} to {a.get('to') or 'today'}"


def _branch_label(user):
    b = scope_branch(user)
    return "All branches" if b is None else BRANCH_NAME.get(b, "")


def _num(v):
    return float(v) if v is not None else None


def _members(user):
    ms = apply_scope(Member.query, Member, user).order_by(Member.branch_id, Member.first_name, Member.last_name).all()
    deps = {}
    for dm in DepartmentMember.query.all():
        deps.setdefault(dm.member_id, []).append(dm.department_id)
    names = {x.id: x.name for x in Department.query.all()}
    rows = [
        [
            m.full_name,
            "Female" if m.gender == "F" else "Male",
            m.dob,
            m.age,
            GROUP_LABEL.get(m.group, ""),
            m.phone,
            m.email,
            m.residence,
            m.marital_status,
            m.home_church.name if m.home_church else "",
            ", ".join(names[i] for i in deps.get(m.id, []) if i in names),
            BRANCH_NAME.get(m.branch_id),
            m.membership_status,
            m.joined_on,
        ]
        for m in ms
    ]
    return "Members", ["Full name", "Gender", "Date of birth", "Age", "Group", "Phone", "Email", "Where they stay", "Marital status", "Home church", "Departments", "Branch", "Status", "Joined"], rows, (), False


def _tithe(user):
    rows = []
    total = 0
    for t in _dated(apply_scope(Tithe.query, Tithe, user), Tithe).order_by(Tithe.date, Tithe.id):
        rows.append([t.date, t.member.full_name, BRANCH_NAME.get(t.branch_id), METHOD.get(t.method, t.method), t.reference or "", _num(t.amount)])
        total += t.amount
    rows.append(["TOTAL", f"{len(rows)} entries", "", "", "", _num(total)])
    return "Tithe", ["Date", "Member", "Branch", "Method", "M-Pesa code / reference", "Amount (KSh)"], rows, (6,), True


def _tithe_by_member(user):
    sums = {}
    for t in _dated(apply_scope(Tithe.query, Tithe, user), Tithe):
        k = t.member_id
        e = sums.setdefault(k, [t.member.full_name, BRANCH_NAME.get(t.branch_id), 0, None, 0])
        e[2] += 1
        e[3] = max(e[3], t.date) if e[3] else t.date
        e[4] += t.amount
    rows = sorted(([a, b, c, d_, _num(e)] for a, b, c, d_, e in sums.values()), key=lambda r: -r[4])
    rows.append(["TOTAL", "", sum(r[2] for r in rows), "", sum(r[4] for r in rows)])
    return "Tithe by member", ["Member", "Branch", "Times given", "Last given", "Total (KSh)"], rows, (5,), True


def _offering(user):
    rows = []
    totals = [0, 0, 0, 0]
    for o in _dated(apply_scope(Offering.query, Offering, user), Offering).order_by(Offering.date, Offering.branch_id):
        counts = o.counts or {}
        rows.append([o.date, o.date.strftime("%A"), o.service, BRANCH_NAME.get(o.branch_id), *[int(counts.get(str(v), 0)) for v in DENOMINATIONS], _num(o.cash_total), _num(o.mpesa_total), _num(o.bank_total), _num(o.total), o.counted_by or ""])
        totals = [totals[0] + o.cash_total, totals[1] + o.mpesa_total, totals[2] + o.bank_total, totals[3] + o.total]
    n = len(DENOMINATIONS)
    rows.append(["TOTAL", "", f"{len(rows)} services", "", *[""] * n, *[_num(x) for x in totals], ""])
    first = 5 + n
    return "Offering", ["Date", "Day", "Service", "Branch", *[f"{v}s" for v in DENOMINATIONS], "Cash (KSh)", "M-Pesa (KSh)", "Bank (KSh)", "Total (KSh)", "Counted by"], rows, tuple(range(first, first + 4)), True


def _pledges(user):
    rows = []
    for p in _dated(apply_scope(Pledge.query, Pledge, user), Pledge, "pledged_on").order_by(Pledge.pledged_on):
        rows.append([p.member.full_name, BRANCH_NAME.get(p.branch_id), p.project, p.pledged_on, p.due_date, _num(p.amount), _num(p.paid), _num(p.balance), p.status])
    rows.append(["TOTAL", "", "", "", "", sum(r[5] for r in rows), sum(r[6] for r in rows), sum(r[7] for r in rows), ""])
    return "Pledges", ["Member", "Branch", "Project", "Pledged on", "Due by", "Pledged (KSh)", "Paid (KSh)", "Balance (KSh)", "Status"], rows, (6, 7, 8), True


def _building_fund(user):
    rows = []
    for f in _dated(apply_scope(FundContribution.query, FundContribution, user), FundContribution).order_by(FundContribution.date):
        rows.append([f.date, f.contributor, BRANCH_NAME.get(f.branch_id), f.project, METHOD.get(f.method, f.method), f.reference or "", _num(f.amount)])
    rows.append(["TOTAL", f"{len(rows)} gifts", "", "", "", "", sum(r[6] for r in rows)])
    return "Building fund", ["Date", "Given by", "Branch", "Project", "Method", "Reference", "Amount (KSh)"], rows, (7,), True


def _inventory(user):
    rows = [
        [i.category, i.name, i.brand, i.model, i.serial_no, i.quantity, i.condition, i.location, i.custodian.full_name if i.custodian else "", i.last_checked, BRANCH_NAME.get(i.branch_id), i.notes or ""]
        for i in apply_scope(InventoryItem.query, InventoryItem, user).order_by(InventoryItem.branch_id, InventoryItem.category, InventoryItem.name)
    ]
    return "Inventory", ["Category", "Item", "Brand", "Model", "Serial no.", "Quantity", "Condition", "Location", "In charge", "Last checked", "Branch", "Notes"], rows, (), False


def _home_church(user):
    rows = []
    for r in _dated(apply_scope(CellReport.query, CellReport, user), CellReport).order_by(CellReport.date, CellReport.cell_id):
        adults, children = len(r.adults or []), len(r.children or [])
        rows.append([r.date, r.cell.name, BRANCH_NAME.get(r.branch_id), r.venue, r.preacher, r.preaching_from, r.worship_leader, adults, children, r.visitors, adults + children + (r.visitors or 0), _num(r.offering), r.signed_by])
    rows.append(["TOTAL", f"{len(rows)} reports", "", "", "", "", "", sum(r[7] for r in rows), sum(r[8] for r in rows), sum(r[9] for r in rows), sum(r[10] for r in rows), sum(r[11] for r in rows), ""])
    return "Home church reports", ["Date", "Home church", "Branch", "Venue", "Preacher", "Preaching from", "Worship leader", "Adults", "Children", "Visitors", "Total present", "Offering (KSh)", "Signed by"], rows, (12,), True


def _home_churches(user):
    q = HomeChurch.query.filter_by(active=True)
    b = scope_branch(user)
    if b is not None:
        q = q.filter(HomeChurch.branch_id == b)
    rows = []
    for c in q.order_by(HomeChurch.branch_id, HomeChurch.name):
        leader = db.session.get(Member, c.leader_member_id) if c.leader_member_id else None
        asst = db.session.get(Member, c.assistant_member_id) if c.assistant_member_id else None
        rows.append([c.name, BRANCH_NAME.get(c.branch_id), c.area, c.venue, c.meeting_time, leader.full_name if leader else "", leader.phone if leader else "", asst.full_name if asst else "", Member.query.filter_by(home_church_id=c.id).count()])
    return "Home churches", ["Home church", "Branch", "Area", "Venue", "Time", "Leader", "Leader's phone", "Assistant", "Members"], rows, (), False


def _departments(user):
    rows = []
    for dep in apply_scope(Department.query, Department, user).order_by(Department.branch_id, Department.name):
        last = DepartmentReport.query.filter_by(department_id=dep.id).order_by(DepartmentReport.date.desc()).first()
        rows.append([dep.name, BRANCH_NAME.get(dep.branch_id), dep.leader.full_name if dep.leader else "", dep.leader.phone if dep.leader else "", dep.assistant.full_name if dep.assistant else "", DepartmentMember.query.filter_by(department_id=dep.id).count(), DepartmentReport.query.filter_by(department_id=dep.id).count(), last.date if last else None, "Active" if dep.active else "Closed"])
    return "Departments", ["Department", "Branch", "Leader", "Leader's phone", "Assistant", "Members", "Reports", "Last report", "Status"], rows, (), False


def _department_reports(user):
    names = {x.id: x.name for x in Department.query.all()}
    rows = [
        [r.date, names.get(r.department_id, ""), BRANCH_NAME.get(r.branch_id), r.title, r.details or "", r.people_involved, r.submitted_by_name or ""]
        for r in _dated(apply_scope(DepartmentReport.query, DepartmentReport, user), DepartmentReport).order_by(DepartmentReport.date, DepartmentReport.department_id)
    ]
    return "Department reports", ["Date", "Department", "Branch", "Title", "Details", "People involved", "Submitted by"], rows, (), False


def _leaders(user):
    q = Leader.query.filter_by(active=True)
    b = scope_branch(user)
    if b is not None:
        q = q.filter((Leader.branch_id == b) | (Leader.role == "Bishop"))
    rows = [[l.role, l.member.full_name, l.member.phone, l.group or "", BRANCH_NAME.get(l.branch_id), l.since] for l in q.order_by(Leader.branch_id, Leader.group, Leader.role)]
    return "Leadership", ["Role", "Name", "Phone", "Group", "Branch", "Since"], rows, (), False


REPORTS = {
    "members": _members,
    "tithe": _tithe,
    "tithe-by-member": _tithe_by_member,
    "offering": _offering,
    "pledges": _pledges,
    "building-fund": _building_fund,
    "inventory": _inventory,
    "home-church-reports": _home_church,
    "home-churches": _home_churches,
    "departments": _departments,
    "department-reports": _department_reports,
    "leadership": _leaders,
}

DATED = {"tithe", "tithe-by-member", "offering", "pledges", "building-fund", "home-church-reports", "department-reports"}


def _sheet(ws, title, headers, rows, money_cols, has_total, period, branch):
    ws.title = title[:31]
    ws.append(["Lord's Altar Ministries International — " + title])
    ws["A1"].font = Font(bold=True, size=13, color="2D3191")
    ws.append([f"{branch}  ·  {period}  ·  generated {date.today():%d %b %Y}"])
    ws["A2"].font = Font(italic=True, color="6B6B6B")
    ws.append([])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    last = ws.max_row
    for row in ws.iter_rows(min_row=5, max_row=last):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "dd mmm yyyy"
    for idx in money_cols:
        for row in ws.iter_rows(min_row=5, max_row=last, min_col=idx, max_col=idx):
            for cell in row:
                cell.number_format = "#,##0"
    if has_total and rows:
        for cell in ws[last]:
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(f"{r[i - 1]:%d %b %Y}") if isinstance(r[i - 1], date) else len(str(r[i - 1])) for r in rows if r[i - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 45)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(last, 4)}"


def _send(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@api.get("/reports/download/<kind>.xlsx")
@login_required(*OFFICE)
def download_report(user, kind):
    if kind == "everything":
        wb = Workbook()
        wb.remove(wb.active)
        for kind, fn in REPORTS.items():
            _sheet(wb.create_sheet(), *fn(user), _period() if kind in DATED else "Current list", _branch_label(user))
        audit(user, "Downloaded report", "Everything")
        db.session.commit()
        return _send(wb, f"LAIM-everything-{date.today():%Y-%m-%d}.xlsx")
    fn = REPORTS.get(kind)
    if not fn:
        raise ApiError(404, "Unknown report.")
    wb = Workbook()
    title, headers, rows, money_cols, has_total = fn(user)
    _sheet(wb.active, title, headers, rows, money_cols, has_total, _period() if kind in DATED else "Current list", _branch_label(user))
    audit(user, "Downloaded report", title)
    db.session.commit()
    return _send(wb, f"LAIM-{kind}-{date.today():%Y-%m-%d}.xlsx")
