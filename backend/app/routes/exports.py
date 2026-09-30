import io
from datetime import date

from flask import request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from ..constants import BRANCHES
from ..extensions import db
from ..models import Member, Tithe
from ..pdfkit import ALTAR, INK, LINE, PAGE_H, PAGE_W, Sheet
from ..security import apply_scope, ensure_can_read, login_required
from . import api
from .giving import tithe_report_data
from .people import OFFICE

BRANCH_NAME = {i: n for i, _, n in BRANCHES}
GROUP_LABEL = {"sunday_school": "Sunday School", "teens": "Teenagers", "junior_youth": "Junior Youth", "senior_youth": "Senior Youth", "fathers": "Fathers", "mothers": "Mothers"}
HEADER_FILL = PatternFill("solid", fgColor="2D3191")


def _xlsx(title, headers, rows, filename, money_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(["Lord's Altar Ministries International — " + title])
    ws["A1"].font = Font(bold=True, size=13, color="2D3191")
    ws.append([f"Generated {date.today():%d %b %Y}"])
    ws.append([])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for idx in money_cols:
        for row in ws.iter_rows(min_row=5, min_col=idx, max_col=idx):
            for cell in row:
                cell.number_format = '#,##0'
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@api.get("/exports/members.xlsx")
@login_required(*OFFICE)
def export_members(user):
    ms = apply_scope(Member.query, Member, user).order_by(Member.branch_id, Member.first_name).all()
    rows = [[m.full_name, "Female" if m.gender == "F" else "Male", m.age, GROUP_LABEL.get(m.group, ""), m.phone, m.email, m.residence, m.marital_status, m.home_church.name if m.home_church else "", BRANCH_NAME.get(m.branch_id), m.membership_status] for m in ms]
    return _xlsx("Members", ["Full name", "Gender", "Age", "Group", "Phone", "Email", "Residence", "Marital status", "Home church", "Branch", "Status"], rows, "LAIM-members.xlsx")


@api.get("/exports/tithe.xlsx")
@login_required(*OFFICE)
def export_tithe(user):
    data = tithe_report_data(user)
    rows = [[r["name"], BRANCH_NAME.get(r["branch_id"]), r["count"], r["last_date"], r["total"]] for r in data["by_member"]]
    rows.append(["TOTAL", "", data["count"], "", data["total"]])
    period = f"{request.args.get('from') or 'start'} to {request.args.get('to') or 'today'}"
    return _xlsx(f"Tithe by member ({period})", ["Member", "Branch", "Times given", "Last given", "Total (KSh)"], rows, "LAIM-tithe.xlsx", money_cols=(5,))


@api.get("/exports/tithe-statement/<int:mid>.pdf")
@login_required(*OFFICE)
def tithe_statement(user, mid):
    m = db.get_or_404(Member, mid)
    ensure_can_read(user, m.branch_id)
    q = Tithe.query.filter_by(member_id=m.id)
    if request.args.get("from"):
        q = q.filter(Tithe.date >= request.args["from"])
    if request.args.get("to"):
        q = q.filter(Tithe.date <= request.args["to"])
    rows = q.order_by(Tithe.date).all()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(PAGE_W, PAGE_H))
    c.setTitle(f"Tithe statement — {m.full_name}")
    s = Sheet(c, "Tithe Statement", f"{m.full_name}  ·  {BRANCH_NAME.get(m.branch_id)}  ·  printed {date.today():%d %b %Y}")
    cols = [s.x0, s.x0 + 45 * mm, s.x0 + 85 * mm]
    right = s.x1
    c.setFont("Manrope-Bold", 8)
    c.setFillColor(INK)
    for x, h in zip(cols, ["DATE", "METHOD", "REFERENCE"]):
        c.drawString(x, s.y, h)
    c.drawRightString(right, s.y, "AMOUNT (KSH)")
    s.y -= 3 * mm
    total = 0
    for t in rows:
        if s.y < 30 * mm:
            s.ensure(40)
        c.setStrokeColor(LINE)
        c.line(s.x0, s.y, s.x1, s.y)
        s.y -= 5.5 * mm
        c.setFont("Manrope", 9.5)
        c.setFillColor(INK)
        c.drawString(cols[0], s.y, f"{t.date:%a %d %b %Y}")
        c.drawString(cols[1], s.y, {"mpesa": "M-Pesa", "bank": "Bank", "cash": "Cash"}[t.method])
        c.drawString(cols[2], s.y, t.reference or "—")
        c.drawRightString(right, s.y, f"{t.amount:,.0f}")
        total += t.amount
        s.y -= 2.5 * mm
    c.setStrokeColor(ALTAR)
    c.setLineWidth(1.2)
    c.line(s.x0, s.y, s.x1, s.y)
    s.y -= 7 * mm
    c.setFont("Sora-Bold", 12)
    c.setFillColor(ALTAR)
    c.drawString(s.x0, s.y, f"Total: {len(rows)} entries")
    c.drawRightString(right, s.y, f"KSh {total:,.0f}")
    c.showPage()
    c.save()
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Tithe-{m.first_name}-{m.last_name}.pdf", mimetype="application/pdf")
